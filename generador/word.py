"""
Word (.docx) con los mismos datos, colores y orden que el PDF.

Solo .docx: nunca con macros. Todo el texto entra como texto plano de Word,
así que no hay forma de que un campo del usuario se interprete como código.
"""
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

FUENTE = "Arial"
ANCHO = 174  # mm útiles en A4 con márgenes de 18 mm


# ---------------------------------------------------------------- utilidades

def rgb(h):
    return RGBColor.from_string(h.lstrip("#").upper())


def run(par, txt, size=9.5, bold=False, color=None, caps=False, espaciado=None):
    r = par.add_run(txt)
    r.font.name = FUENTE
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.all_caps = caps
    if color:
        r.font.color.rgb = rgb(color)
    if espaciado is not None:
        rpr = r._element.get_or_add_rPr()
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:val"), str(espaciado))
        rpr.append(sp)
    return r


def parrafo(contenedor, txt="", size=9.5, bold=False, color=None, alinear=None,
            antes=0, despues=3, caps=False, espaciado=None):
    par = contenedor.add_paragraph()
    pf = par.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(antes), Pt(despues), 1.15
    if alinear:
        par.alignment = alinear
    if txt:
        lineas = txt.split("\n")
        for i, linea in enumerate(lineas):
            r = run(par, linea, size, bold, color, caps, espaciado)
            if i < len(lineas) - 1:
                r.add_break()
    return par


def primer_parrafo(celda, **kw):
    """Usa el párrafo vacío que trae cada celda en vez de dejar uno de sobra."""
    par = celda.paragraphs[0]
    pf = par.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(kw.get("antes", 0)), Pt(kw.get("despues", 3)), 1.15
    if kw.get("alinear"):
        par.alignment = kw["alinear"]
    return par


def rotulo(celda_o_doc, txt, colores, primero=False):
    if primero:
        par = primer_parrafo(celda_o_doc, despues=2)
    else:
        par = parrafo(celda_o_doc, despues=2)
    run(par, txt, 7, True, colores["acento_texto"], caps=True, espaciado=20)
    return par


def sombrear(celda, color):
    tcpr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color.lstrip("#").upper())
    tcpr.append(shd)


def bordes(tabla, color="E3E5E8", lados=("insideH", "bottom"), grosor=6):
    tblpr = tabla._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{lado}")
        if lado in lados:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(grosor))
            el.set(qn("w:color"), color.lstrip("#").upper())
        else:
            el.set(qn("w:val"), "nil")
        b.append(el)
    tblpr.append(b)


def borde_celda(celda, color, lados=("top", "left", "bottom", "right"), grosor=6):
    tcpr = celda._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for lado in lados:
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(grosor))
        el.set(qn("w:color"), color.lstrip("#").upper())
        b.append(el)
    tcpr.append(b)


def margen_celdas(tabla, arriba=60, abajo=60, izq=90, der=90):
    tblpr = tabla._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for lado, val in (("top", arriba), ("bottom", abajo), ("left", izq), ("right", der)):
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblpr.append(mar)


def tabla(contenedor, filas, anchos_mm, sin_bordes=True):
    t = contenedor.add_table(rows=filas, cols=len(anchos_mm))
    t.autofit = False
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    t._tbl.tblPr.append(layout)
    for fila in t.rows:
        for celda, ancho in zip(fila.cells, anchos_mm):
            celda.width = Mm(ancho)
    if sin_bordes:
        bordes(t, lados=())
    return t


def no_partir_fila(fila, *textos):
    """Evita que una fila corta quede partida entre dos páginas. Una fila larga
    sí se deja partir: si no, Word corta en silencio lo que no cabe en la hoja."""
    if sum(len(t or "") + 60 * (t or "").count("\n") for t in textos) > 500:
        return
    trpr = fila._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    trpr.append(el)


def repetir_encabezado(fila):
    trpr = fila._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    trpr.append(el)


def campo(par, instruccion, size, color):
    """Inserta un campo de Word (PAGE, NUMPAGES)."""
    for tipo, valor in (("begin", None), ("instr", instruccion), ("separate", None), ("end", None)):
        r = run(par, "", size, color=color)
        if tipo == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {valor} "
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
        r._element.append(el)
        if tipo == "separate":
            run(par, "1", size, color=color)


def espacio(doc, pt):
    parrafo(doc, despues=0, antes=0).paragraph_format.space_after = Pt(pt)


# ------------------------------------------------------------------- bloques

def configurar(doc, v):
    normal = doc.styles["Normal"]
    normal.font.name = FUENTE
    normal.font.size = Pt(9.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
    s = doc.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.left_margin = s.right_margin = Mm(18)
    s.top_margin, s.bottom_margin = Mm(16), Mm(16)
    s.footer_distance = Mm(8)
    pie(s, v)


def pie(seccion, v):
    # Tabla y no tabulación: el estilo de pie de Word trae tabulaciones propias
    # que mandaban el número de página al centro.
    c = v["colores"]
    t = seccion.footer.add_table(rows=1, cols=2, width=Mm(ANCHO))
    t.autofit = False
    bordes(t, lados=())
    izq, der = t.rows[0].cells
    izq.width = der.width = Mm(ANCHO / 2)
    run(primer_parrafo(izq, despues=0), v["codigo"], 7.5, True, c["tinta_3"])
    par = primer_parrafo(der, despues=0, alinear=WD_ALIGN_PARAGRAPH.RIGHT)
    run(par, "Página ", 7.5, color=c["tinta_3"])
    campo(par, "PAGE", 7.5, c["tinta_3"])
    run(par, " de ", 7.5, color=c["tinta_3"])
    campo(par, "NUMPAGES", 7.5, c["tinta_3"])


def encabezado(doc, v, ruta_logo):
    c, n = v["colores"], v["negocio"]
    t = tabla(doc, 1, [80, ANCHO - 80])
    izq, der = t.rows[0].cells
    izq.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    par = primer_parrafo(izq, despues=0)
    if ruta_logo:
        par.add_run().add_picture(str(ruta_logo), height=Mm(15))
    else:
        run(par, n["nombre"], 17, True, c["principal_texto"])

    par = primer_parrafo(der, despues=0, alinear=WD_ALIGN_PARAGRAPH.RIGHT)
    lineas = []
    if n["id_valor"]:
        lineas.append(f"{n['id_etiqueta']} {n['id_valor']}")
    if n["direccion"]:
        lineas.append(n["direccion"])
    contacto = " · ".join(x for x in (n["email"], n["telefono"]) if x)
    if contacto:
        lineas.append(contacto)
    if n["web"]:
        lineas.append(n["web"])
    if ruta_logo:
        run(par, n["nombre"], 9.5, True, c["tinta"])
    for i, linea in enumerate(lineas):
        if ruta_logo or i:
            run(par, "", 8, color=c["tinta_2"]).add_break()
        run(par, linea, 8, color=c["tinta_2"])

    franja = parrafo(doc, antes=6, despues=10)
    ppr = franja._p.get_or_add_pPr()
    borde = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "18")
    bottom.set(qn("w:color"), c["principal"].lstrip("#").upper())
    borde.append(bottom)
    ppr.append(borde)


def tabla_items(doc, v):
    c = v["colores"]
    desc = v["hay_desc_linea"]
    cols = ["#", "Descripción", "Cant.", "Precio unit."] + (["Desc."] if desc else []) + ["Total"]
    fijos = [8, 20, 28] + ([16] if desc else []) + [28]
    anchos = [fijos[0], ANCHO - sum(fijos)] + fijos[1:]
    t = tabla(doc, 1 + len(v["lineas"]), anchos, sin_bordes=False)
    bordes(t, c["borde"], lados=("insideH", "bottom"))
    margen_celdas(t, 80, 80, 90, 90)

    cab = t.rows[0]
    repetir_encabezado(cab)
    for i, (celda, titulo) in enumerate(zip(cab.cells, cols)):
        sombrear(celda, c["principal"])
        par = primer_parrafo(celda, despues=0,
                             alinear=WD_ALIGN_PARAGRAPH.RIGHT if i >= 2 else None)
        run(par, titulo, 7.5, True, c["sobre_principal"], caps=True)

    for fila, it in zip(t.rows[1:], v["lineas"]):
        no_partir_fila(fila, it["descripcion"], it["detalle"])
        valores = [str(it["n"]), None, it["cantidad_txt"] + (f" {it['unidad']}" if it["unidad"] else ""),
                   it["precio_txt"]] + ([it["descuento_txt"]] if desc else []) + [it["total_txt"]]
        for i, (celda, valor) in enumerate(zip(fila.cells, valores)):
            celda.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            if i == 1:
                par = primer_parrafo(celda, despues=0)
                for j, linea in enumerate(it["descripcion"].split("\n")):
                    r = run(par, linea, 9.5, True, c["tinta"])
                    if j < len(it["descripcion"].split("\n")) - 1:
                        r.add_break()
                if it["detalle"]:
                    parrafo(celda, it["detalle"], 8, color=c["tinta_3"], despues=0, antes=1)
            else:
                par = primer_parrafo(celda, despues=0,
                                     alinear=WD_ALIGN_PARAGRAPH.RIGHT if i >= 2 else None)
                run(par, valor, 8.5 if i == 0 else 9.5, color=c["tinta_3"] if i == 0 else c["tinta"])


def condiciones_y_totales(doc, v):
    c = v["colores"]
    espacio(doc, 6)
    t = tabla(doc, 1, [ANCHO - 76, 76])
    izq, der = t.rows[0].cells

    rotulo(izq, "Validez", c, primero=True)
    parrafo(izq, f"Esta oferta es válida hasta el {v['vence']}.", 8.8, color=c["tinta_2"], despues=8)
    for etiqueta, texto in v["condiciones"]:
        rotulo(izq, etiqueta, c)
        parrafo(izq, texto, 8.8, color=c["tinta_2"], despues=8)

    primer_parrafo(der, despues=0)
    tt = der.add_table(rows=len(v["totales"]), cols=2)
    tt.autofit = False
    margen_celdas(tt, 70, 70, 110, 110)
    bordes(tt, lados=())
    for fila, tot in zip(tt.rows, v["totales"]):
        a, b = fila.cells
        a.width, b.width = Mm(38), Mm(36)
        if tot["clave"] == "total":
            # Una sola celda: dos celdas pintadas dejan una línea fina entre ellas.
            celda = a.merge(b)
            sombrear(celda, c["principal"])
            par = primer_parrafo(celda, despues=0)
            par.paragraph_format.tab_stops.add_tab_stop(Mm(70), WD_TAB_ALIGNMENT.RIGHT)
            run(par, tot["etiqueta"], 11, True, c["sobre_principal"])
            run(par, "\t" + tot["texto"], 11, True, c["sobre_principal"])
            continue
        borde_celda(a, c["borde"], ("bottom",))
        borde_celda(b, c["borde"], ("bottom",))
        run(primer_parrafo(a, despues=0), tot["etiqueta"], 9, color=c["tinta_2"])
        run(primer_parrafo(b, despues=0, alinear=WD_ALIGN_PARAGRAPH.RIGHT), tot["texto"], 9, color=c["tinta"])
    if v["nota_impuesto"]:
        parrafo(der, v["nota_impuesto"], 7.5, color=c["tinta_3"], alinear=WD_ALIGN_PARAGRAPH.RIGHT, antes=3)


def aceptacion(doc, v):
    c = v["colores"]
    espacio(doc, 26)
    t = tabla(doc, 1, [80, 14, 80])
    no_partir_fila(t.rows[0])
    for celda, texto in ((t.rows[0].cells[0], "Nombre y firma de aceptación"),
                         (t.rows[0].cells[2], "Fecha")):
        borde_celda(celda, c["tinta_3"], ("top",))
        run(primer_parrafo(celda, antes=2, despues=0), texto, 7.5, color=c["tinta_3"])
    parrafo(doc, v["aviso"], 7.5, color=c["tinta_3"], alinear=WD_ALIGN_PARAGRAPH.CENTER, antes=18)


# --------------------------------------------------------------- documentos

def cotizacion(doc, v, ruta_logo):
    c, cl = v["colores"], v["cliente"]
    encabezado(doc, v, ruta_logo)

    t = tabla(doc, 1, [ANCHO - 90, 90])
    izq, der = t.rows[0].cells
    der.vertical_alignment = WD_ALIGN_VERTICAL.BOTTOM
    run(primer_parrafo(izq, despues=0), v["nombre_documento"], 24, True, c["tinta"])
    run(parrafo(izq, despues=0), f"N° {v['codigo']}", 9, True, c["acento_texto"])
    par = primer_parrafo(der, despues=0, alinear=WD_ALIGN_PARAGRAPH.RIGHT)
    run(par, "EMISIÓN  ", 7, True, c["acento_texto"], espaciado=20)
    run(par, v["emision"], 9)
    par = parrafo(der, despues=0, alinear=WD_ALIGN_PARAGRAPH.RIGHT)
    run(par, "VÁLIDA HASTA  ", 7, True, c["acento_texto"], espaciado=20)
    run(par, v["vence"], 9)

    espacio(doc, 8)
    t = tabla(doc, 1, [ANCHO - 66, 4, 62])
    margen_celdas(t, 130, 130, 160, 160)
    cli, _, tot = t.rows[0].cells
    borde_celda(cli, c["borde"])
    rotulo(cli, "Para", c, primero=True)
    run(parrafo(cli, despues=1), cl["nombre"], 11, True, c["tinta"])
    datos = [cl.get("empresa")]
    if cl.get("id_tributario"):
        datos.append(f"{v['negocio']['id_etiqueta'] or 'ID'} {cl['id_tributario']}")
    if cl.get("contacto"):
        datos.append(f"Atención: {cl['contacto']}")
    datos.append(" · ".join(x for x in (cl.get("email"), cl.get("telefono")) if x))
    datos.append(cl.get("direccion"))
    datos = [d for d in datos if d]
    if datos:
        parrafo(cli, "\n".join(datos), 8.5, color=c["tinta_2"], despues=0)

    sombrear(tot, c["tinte"])
    tot.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    rotulo(tot, "Total", c, primero=True)
    run(parrafo(tot, despues=0), v["total_txt"], 17, True, c["principal_texto"])
    if v["nota_total"]:
        parrafo(tot, v["nota_total"], 7.5, color=c["tinta_3"], despues=0)

    espacio(doc, 8)
    tabla_items(doc, v)
    condiciones_y_totales(doc, v)
    aceptacion(doc, v)


def seccion(doc, numero, titulo, colores):
    par = parrafo(doc, antes=16, despues=0)
    par.paragraph_format.keep_with_next = True
    run(par, f"{numero:02d}", 20, True, colores["acento_texto"])
    par = parrafo(doc, titulo, 15, True, colores["tinta"], despues=6)
    par.paragraph_format.keep_with_next = True


def parrafos(doc, lista, colores):
    for p in lista:
        parrafo(doc, p, 9.8, color=colores["tinta_2"], despues=6)


def propuesta(doc, v, ruta_logo):
    c, p, n = v["colores"], v["propuesta"], v["negocio"]
    doc.sections[0].different_first_page_header_footer = True

    portada = tabla(doc, 1, [ANCHO])
    fila = portada.rows[0]
    fila.height, fila.height_rule = Mm(255), WD_ROW_HEIGHT_RULE.EXACTLY
    celda = fila.cells[0]
    sombrear(celda, c["principal"])
    margen_celdas(portada, 900, 700, 800, 800)
    sobre = c["sobre_principal"]

    par = primer_parrafo(celda, despues=0)
    if ruta_logo:
        # La tarjeta blanca asegura que el logo se vea sobre cualquier color.
        tl = celda.add_table(rows=1, cols=1)
        margen_celdas(tl, 120, 120, 160, 160)
        bordes(tl, lados=())
        sombrear(tl.rows[0].cells[0], "#ffffff")
        tl.rows[0].cells[0].width = Mm(62)
        primer_parrafo(tl.rows[0].cells[0], despues=0).add_run().add_picture(str(ruta_logo), height=Mm(13))
    else:
        run(par, n["nombre"], 17, True, sobre)

    parrafo(celda, f"PROPUESTA · {v['codigo']}", 8.5, True, sobre, antes=150, despues=14, espaciado=40)
    parrafo(celda, p["titulo"], 32, True, sobre, despues=14)
    par = parrafo(celda, despues=0)
    run(par, "Preparada para ", 13, color=sobre)
    run(par, v["cliente"].get("empresa") or v["cliente"]["nombre"], 13, True, sobre)

    parrafo(celda, n["nombre"], 11, True, sobre, antes=170, despues=1)
    contacto = " · ".join(x for x in (n["email"], n["telefono"]) if x)
    if contacto:
        parrafo(celda, contacto, 9, color=sobre, despues=1)
    parrafo(celda, f"{v['emision']} · Válida hasta el {v['vence']}", 9, color=sobre, despues=0)

    doc.add_page_break()
    encabezado(doc, v, ruta_logo)

    k = 0
    if p["contexto"]:
        k += 1
        seccion(doc, k, "Contexto", c)
        parrafos(doc, p["contexto"], c)
    if p["solucion"]:
        k += 1
        seccion(doc, k, "Nuestra propuesta", c)
        parrafos(doc, p["solucion"], c)
    if p["incluye"] or p["no_incluye"]:
        k += 1
        seccion(doc, k, "Alcance", c)
        t = tabla(doc, 1, [85, 4, 85])
        margen_celdas(t, 130, 130, 160, 160)
        no_partir_fila(t.rows[0], *p["incluye"], *p["no_incluye"])
        si, _, no = t.rows[0].cells
        for celda, titulo, lista, fondo in ((si, "Incluye", p["incluye"], True),
                                            (no, "No incluye", p["no_incluye"], False)):
            if not lista:
                continue
            if fondo:
                sombrear(celda, c["tinte"])
            else:
                borde_celda(celda, c["borde"])
            rotulo(celda, titulo, c, primero=True)
            for x in lista:
                par = parrafo(celda, despues=3)
                run(par, "•  " if fondo else "–  ", 9.3, True, c["principal"] if fondo else c["tinta_3"])
                run(par, x, 9.3, color=c["tinta_2"])
    if p["etapas"]:
        k += 1
        seccion(doc, k, "Etapas y plazos", c)
        t = tabla(doc, len(p["etapas"]), [ANCHO - 40, 40], sin_bordes=False)
        bordes(t, c["borde"], lados=("insideH",))
        margen_celdas(t, 90, 90, 0, 0)
        for fila, et in zip(t.rows, p["etapas"]):
            no_partir_fila(fila, et["nombre"], et.get("descripcion"))
            a, b = fila.cells
            run(primer_parrafo(a, despues=1), et["nombre"], 10, True, c["tinta"])
            if et.get("descripcion"):
                parrafo(a, et["descripcion"], 9.3, color=c["tinta_2"], despues=0)
            run(primer_parrafo(b, despues=0, alinear=WD_ALIGN_PARAGRAPH.RIGHT),
                et.get("duracion", ""), 8.5, True, c["acento_texto"])

    k += 1
    seccion(doc, k, "Inversión", c)
    tabla_items(doc, v)
    condiciones_y_totales(doc, v)

    if p["por_que_nosotros"]:
        k += 1
        seccion(doc, k, "Por qué nosotros", c)
        parrafos(doc, p["por_que_nosotros"], c)
    k += 1
    seccion(doc, k, "Próximos pasos", c)
    parrafos(doc, p["proximos_pasos"] or
             ["Para avanzar, basta con responder confirmando la aceptación de esta propuesta."], c)
    aceptacion(doc, v)


def generar(vista, ruta_logo, destino):
    doc = Document()
    configurar(doc, vista)
    # Quita el párrafo vacío inicial que trae el documento nuevo.
    cuerpo = doc.element.body
    for p in list(cuerpo.iterchildren(qn("w:p"))):
        cuerpo.remove(p)
    if vista["tipo"] == "propuesta":
        propuesta(doc, vista, ruta_logo)
    else:
        cotizacion(doc, vista, ruta_logo)
    doc.core_properties.title = f"{vista['nombre_documento']} {vista['codigo']}"
    doc.core_properties.author = vista["negocio"]["nombre"]
    doc.save(destino)
    return destino
