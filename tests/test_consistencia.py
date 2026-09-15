"""
Los tres formatos salen del mismo cálculo: esta prueba lo comprueba leyendo los
archivos terminados, no la lógica interna.
"""
import shutil
import unittest

import apoyo
from pdf import buscar_navegador


class TresFormatosMismoTotal(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.carpeta = apoyo.carpeta_limpia("consistencia")
        negocio = apoyo.json.loads((apoyo.RAIZ / "ejemplos" / "negocio-ejemplo.json").read_text(encoding="utf-8"))
        apoyo.guardar(cls.carpeta / "mi-negocio" / "negocio.json", negocio)
        shutil.copy(apoyo.RAIZ / "ejemplos" / "logo-ejemplo.png", cls.carpeta / "mi-negocio" / "logo.png")
        cls.resultados = {}
        for tipo, total in (("cotizacion", "$981.036"), ("propuesta", "$1.752.275")):
            r = apoyo.generar(cls.carpeta, "crear", f"ejemplos/{tipo}-ejemplo.json")
            cls.resultados[tipo] = (r, total)

    def archivos(self, tipo):
        prefijo = "COT" if tipo == "cotizacion" else "PROP"
        return next((self.carpeta / "mis-documentos").glob(f"{prefijo}-0001*"))

    def test_se_crean_los_tres(self):
        for tipo, (r, _) in self.resultados.items():
            with self.subTest(tipo):
                esperados = 3 if buscar_navegador() else 2
                self.assertEqual(r.stdout.count("CREADO"), esperados, r.stdout + r.stderr)

    def test_word_muestra_el_mismo_total(self):
        from docx import Document
        for tipo, (_, total) in self.resultados.items():
            with self.subTest(tipo):
                doc = Document(next(self.archivos(tipo).glob("*.docx")))
                texto = []
                def recorrer(tablas):
                    for t in tablas:
                        for fila in t.rows:
                            for celda in fila.cells:
                                texto.append(celda.text)
                                recorrer(celda.tables)
                recorrer(doc.tables)
                self.assertIn(total, "\n".join(texto))

    def test_excel_guarda_el_mismo_total(self):
        from openpyxl import load_workbook
        for tipo, (_, total) in self.resultados.items():
            with self.subTest(tipo):
                hoja = load_workbook(next(self.archivos(tipo).glob("*.xlsx")), data_only=True).worksheets[0]
                valores = [c.value for fila in hoja.iter_rows() for c in fila if isinstance(c.value, (int, float))]
                self.assertIn(int(total.replace("$", "").replace(".", "")), valores)

    @unittest.skipUnless(buscar_navegador(), "requiere Chrome o Edge")
    def test_pdf_muestra_el_mismo_total(self):
        from pypdf import PdfReader
        for tipo, (_, total) in self.resultados.items():
            with self.subTest(tipo):
                lector = PdfReader(next(self.archivos(tipo).glob("*.pdf")))
                self.assertIn(total, "\n".join(p.extract_text() for p in lector.pages))

    def test_aviso_de_que_no_es_factura(self):
        from pypdf import PdfReader
        if not buscar_navegador():
            self.skipTest("requiere Chrome o Edge")
        lector = PdfReader(next(self.archivos("cotizacion").glob("*.pdf")))
        self.assertIn("No constituye factura", "\n".join(p.extract_text() for p in lector.pages))


if __name__ == "__main__":
    unittest.main()
