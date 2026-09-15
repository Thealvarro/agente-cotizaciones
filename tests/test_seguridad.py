"""
Cada prueba corresponde a un punto del brief de seguridad.
"""
import unittest

import apoyo
from casos import negocio, documento
from datos import nombre_seguro, validar_documento, validar_negocio


class TextoDelUsuarioEnElPDF(unittest.TestCase):
    def test_html_pegado_sale_escapado(self):
        from contenido import armar
        import pdf
        doc = documento([{"descripcion": "<img src=x onerror=alert(1)>", "cantidad": 1, "precio_unitario": 10}],
                        cliente={"nombre": "<script>alert('x')</script>"}, folio=1)
        html = pdf.html(armar(doc, negocio()), None)
        self.assertNotIn("<script>alert", html)
        self.assertNotIn("<img src=x", html)
        self.assertIn("&lt;script&gt;", html)

    def test_color_no_puede_inyectar_css(self):
        n = negocio(color_principal="#000;}body{display:none}")
        self.assertTrue(any("color_principal" in e for e in validar_negocio(n)))


class FormulasInyectadasEnExcel(unittest.TestCase):
    def test_texto_que_parece_formula_queda_como_texto(self):
        from openpyxl import load_workbook
        from contenido import armar
        import excel

        peligrosos = ['=HYPERLINK("http://evil.example","clic")', "+cmd|' /C calc'!A0", "-2+3", "@SUM(A1:A9)"]
        items = [{"descripcion": t, "detalle": t, "unidad": "=1+1", "cantidad": 1, "precio_unitario": 1}
                 for t in peligrosos]
        doc = documento(items, cliente={"nombre": "=1+1", "empresa": "@SUM(1)"}, folio=1)
        ruta = apoyo.carpeta_limpia("inyeccion") / "inyeccion.xlsx"
        excel.generar(armar(doc, negocio()), None, ruta)

        hoja = load_workbook(ruta).active
        formulas = [c.value for fila in hoja.iter_rows() for c in fila if c.data_type == "f"]
        textos = [c.value for fila in hoja.iter_rows() for c in fila if c.data_type == "s"]
        for f in formulas:
            # Las únicas fórmulas son las de montos que escribe el generador.
            self.assertRegex(f, r"^=(ROUND|SUM|-ROUND|G\d)", f)
        for t in peligrosos:
            # Siguen ahí, pero como texto inofensivo (descripción y detalle comparten celda).
            self.assertTrue(any(t in texto for texto in textos), t)


class Logo(unittest.TestCase):
    def setUp(self):
        self.carpeta = apoyo.carpeta_limpia("logo")

    def _probar(self, nombre, contenido):
        import importlib, os
        os.environ["COTIZAR_DATOS"] = str(self.carpeta)
        import datos
        importlib.reload(datos)
        (self.carpeta / "mi-negocio" / nombre).write_bytes(contenido)
        try:
            return datos.ruta_logo({"logo": nombre})
        finally:
            del os.environ["COTIZAR_DATOS"]
            importlib.reload(datos)

    def test_rechaza_svg(self):
        # Exception y no DatosInvalidos: _probar recarga el módulo y la clase cambia de identidad.
        with self.assertRaises(Exception) as ctx:
            self._probar("logo.svg", b"<svg onload='alert(1)'></svg>")
        self.assertIn("SVG", str(ctx.exception))

    def test_rechaza_archivo_disfrazado_de_png(self):
        with self.assertRaises(Exception) as ctx:
            self._probar("logo.png", b"<svg onload='alert(1)'></svg>")
        self.assertIn("no es realmente", str(ctx.exception))

    def test_rechaza_salir_de_la_carpeta(self):
        with self.assertRaises(Exception) as ctx:
            self._probar("../../fuera.png", b"\x89PNG\r\n\x1a\n")
        self.assertIn("dentro de la carpeta", str(ctx.exception))

    def test_acepta_png_real(self):
        ruta = self._probar("logo.png", (apoyo.RAIZ / "ejemplos" / "logo-ejemplo.png").read_bytes())
        self.assertEqual(ruta.name, "logo.png")


class NombresDeArchivo(unittest.TestCase):
    def test_no_permite_rutas_ni_nombres_reservados(self):
        for peligroso in ["../../etc/passwd", "..\\..\\Windows", "CON", 'a<b>:"c|d?*', "", "   ...   "]:
            with self.subTest(peligroso):
                seguro = nombre_seguro(peligroso)
                self.assertTrue(seguro)
                self.assertNotIn("..", seguro)
                for caracter in '<>:"/\\|?*':
                    self.assertNotIn(caracter, seguro)
                self.assertNotEqual(seguro.upper(), "CON")


class Datos(unittest.TestCase):
    def test_monto_con_separador_de_miles_se_rechaza(self):
        # "1.500" es mil quinientos en Chile y uno coma cinco en México: no se adivina.
        doc = documento([{"descripcion": "X", "cantidad": 1, "precio_unitario": "1.500"}])
        self.assertTrue(any("sin separadores" in e for e in validar_documento(doc)))

    def test_descuento_fuera_de_rango(self):
        doc = documento([{"descripcion": "X", "cantidad": 1, "precio_unitario": 10, "descuento_pct": 150}])
        self.assertTrue(validar_documento(doc))

    def test_archivo_fuera_del_proyecto_se_rechaza(self):
        carpeta = apoyo.carpeta_limpia("fuera")
        apoyo.guardar(carpeta / "mi-negocio" / "negocio.json", negocio())
        r = apoyo.generar(carpeta, "previa", "C:/Windows/win.ini" if apoyo.sys.platform == "win32" else "/etc/hosts")
        self.assertNotEqual(r.returncode, 0)
        self.assertIn("dentro de la carpeta del proyecto", r.stdout)


class DatosDeClientesNoSeSuben(unittest.TestCase):
    def test_gitignore_protege_negocio_y_documentos(self):
        gitignore = (apoyo.RAIZ / ".gitignore").read_text(encoding="utf-8")
        for carpeta in ("mi-negocio/", "mis-documentos/", ".venv/"):
            self.assertIn(carpeta, gitignore)


class Folios(unittest.TestCase):
    def test_folio_no_se_repite_ni_borrando_el_contador(self):
        carpeta = apoyo.carpeta_limpia("folios")
        apoyo.guardar(carpeta / "mi-negocio" / "negocio.json", negocio())
        borrador = carpeta / "mis-documentos" / "borradores" / "doc.json"
        doc = documento([{"descripcion": "X", "cantidad": 1, "precio_unitario": 1000}])

        folios = []
        for i in range(3):
            if i == 2:
                (carpeta / "mis-documentos" / "folios.json").unlink()   # alguien borra el contador
            apoyo.guardar(borrador, doc)
            r = apoyo.generar(carpeta, "crear", str(borrador.relative_to(apoyo.RAIZ)))
            self.assertIn(r.returncode, (0, 3), r.stdout + r.stderr)
            folios.append(max(int(p.parent.name[4:8]) for p in (carpeta / "mis-documentos").glob("COT-*/documento.json")))
        self.assertEqual(folios, [1, 2, 3])

    def test_regenerar_mantiene_el_folio(self):
        carpeta = apoyo.carpeta_limpia("regenerar")
        apoyo.guardar(carpeta / "mi-negocio" / "negocio.json", negocio())
        borrador = carpeta / "mis-documentos" / "borradores" / "doc.json"
        apoyo.guardar(borrador, documento([{"descripcion": "X", "cantidad": 1, "precio_unitario": 1000}]))
        apoyo.generar(carpeta, "crear", str(borrador.relative_to(apoyo.RAIZ)))
        existente = next((carpeta / "mis-documentos").glob("COT-0001*/documento.json"))
        apoyo.generar(carpeta, "crear", str(existente.relative_to(apoyo.RAIZ)))
        self.assertEqual(len(list((carpeta / "mis-documentos").glob("COT-*"))), 1)


if __name__ == "__main__":
    unittest.main()
